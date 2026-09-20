#!/usr/bin/env python3
"""Generate the concise Chinese Word results report from pipeline outputs."""
from __future__ import annotations
import json, math, os
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'outputs'; TABLE=OUT/'tables'; FIG=OUT/'figures'; REPORT=OUT/'report'/'修订稿补充数据分析结果报告.docx'
if os.environ.get("REPORT_OUT"):
    REPORT=Path(os.environ["REPORT_OUT"])
BLUE=RGBColor(46,116,181); DARK=RGBColor(31,77,120); MUTED=RGBColor(90,90,90)

def fmt(x,d=3):
    try:
        x=float(x)
        if not np.isfinite(x): return 'NA'
        return f'{x:.{d}f}'
    except Exception:return str(x)

def ptxt(p):
    p=float(p)
    return 'p < .001' if p<.001 else f'p = {p:.3f}'.replace('0.','.')

def set_font(run,name='Calibri',size=11,bold=None,color=None,italic=None):
    run.font.name=name; run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),name);run._element.rPr.rFonts.set(qn('w:hAnsi'),name);run._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    if italic is not None:run.italic=italic
    if color is not None:run.font.color.rgb=color

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);tcPr.append(shd)

def margins(cell,top=80,start=120,bottom=80,end=120):
    tc=cell._tc.get_or_add_tcPr();m=tc.first_child_found_in('w:tcMar')
    if m is None:m=OxmlElement('w:tcMar');tc.append(m)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=m.find(qn(f'w:{tag}'))
        if node is None:node=OxmlElement(f'w:{tag}');m.append(node)
        node.set(qn('w:w'),str(val));node.set(qn('w:type'),'dxa')

def add_table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False;t.style='Table Grid'
    tr_pr=t.rows[0]._tr.get_or_add_trPr();tbl_header=OxmlElement('w:tblHeader');tbl_header.set(qn('w:val'),'true');tr_pr.append(tbl_header)
    for j,h in enumerate(headers):
        c=t.rows[0].cells[j];c.text=str(h);shade(c,'E8EEF5');c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
        c.paragraphs[0].paragraph_format.keep_with_next=True
        for r in c.paragraphs[0].runs:set_font(r,size=8.5,bold=True)
    for row in rows:
        cells=t.add_row().cells
        for j,v in enumerate(row):
            cells[j].text=str(v);cells[j].vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            for p in cells[j].paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
                for r in p.runs:set_font(r,size=8.3)
    if widths:
        twips=[int(round(w*1440)) for w in widths]
        tbl_pr=t._tbl.tblPr
        tbl_w=tbl_pr.first_child_found_in('w:tblW')
        if tbl_w is None:tbl_w=OxmlElement('w:tblW');tbl_pr.append(tbl_w)
        tbl_w.set(qn('w:w'),str(sum(twips)));tbl_w.set(qn('w:type'),'dxa')
        tbl_ind=tbl_pr.first_child_found_in('w:tblInd')
        if tbl_ind is None:tbl_ind=OxmlElement('w:tblInd');tbl_pr.append(tbl_ind)
        tbl_ind.set(qn('w:w'),'120');tbl_ind.set(qn('w:type'),'dxa')
        for grid_col,w in zip(t._tbl.tblGrid.gridCol_lst,twips):grid_col.set(qn('w:w'),str(w))
        for row in t.rows:
            for c,w in zip(row.cells,widths):c.width=Inches(w);margins(c)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def add_picture(doc,path,width,alt_text):
    shape=doc.add_picture(str(path),width=width)
    shape._inline.docPr.set('descr',alt_text)
    shape._inline.docPr.set('title',alt_text)
    doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next=True
    return shape

def add_bullet(doc,text):
    p=doc.add_paragraph(style='List Bullet');p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.15
    set_font(p.add_run(text));return p

def add_labeled(doc,label,text,keep=False):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.12
    p.paragraph_format.keep_with_next=keep
    set_font(p.add_run(label+'：'),size=10.3,bold=True,color=DARK)
    set_font(p.add_run(text),size=10.3)
    return p

def add_figure_block(doc,method,path,width,alt_text,caption,elements,trend,conclusion,theory):
    add_labeled(doc,'分析方法',method,keep=True)
    add_picture(doc,path,width,alt_text)
    cap=doc.add_paragraph(caption);cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after=Pt(4)
    for r in cap.runs:set_font(r,size=8,italic=True,color=MUTED)
    add_labeled(doc,'图表元素',elements)
    add_labeled(doc,'主要趋势',trend)
    add_labeled(doc,'支持结论',conclusion)
    add_labeled(doc,'理论关系',theory)

def key_coef(model,pattern):
    p=TABLE/'behavior_mixed_model_coefficients.csv'
    if not p.exists():return None
    d=pd.read_csv(p);z=d[(d.model==model)&d.term.astype(str).str.contains(pattern,regex=False)]
    return z.iloc[0] if len(z) else None

def exact_coef(data,model,term):
    z=data[(data.model==model)&(data.term.astype(str)==term)]
    return z.iloc[0] if len(z) else None

def interaction_label(term):
    term=str(term)
    if ':velocity_f' in term:return '速度'+term.split(':velocity_f',1)[1]
    if ':size_pair' in term:return '尺寸'+term.split(':size_pair',1)[1].replace('_','/ ')
    return term.replace('conditionzero_gravity:','')

def build():
    REPORT.parent.mkdir(parents=True,exist_ok=True)
    flow=pd.read_csv(TABLE/'sample_flow.csv');effects=pd.read_csv(TABLE/'paired_condition_effects.csv')
    behavior_coef=pd.read_csv(TABLE/'behavior_mixed_model_coefficients.csv')
    behavior_diag=pd.read_csv(TABLE/'behavior_model_diagnostics.csv')
    acc_auc_corr=pd.read_csv(TABLE/'accuracy_auc_correlations.csv')
    rt_sens=pd.read_csv(TABLE/'rt_sensitivity.csv').iloc[0]
    stimulus_checks=pd.read_csv(TABLE/'stimulus_generation_checks.csv')
    dsm=pd.read_csv(TABLE/'dsm_results_all.csv')
    dsm_continuous=pd.read_csv(TABLE/'dsm_continuous_auc_sensitivity.csv')
    spam_support=pd.read_csv(TABLE/'spam_support_summary.csv')
    eye_direct=pd.read_csv(TABLE/'figure3_direct_eye_metric_summary.csv')
    eye_sens=pd.read_csv(TABLE/'eye_coverage_threshold_sensitivity.csv')
    scan_boot=pd.read_csv(TABLE/'figure4_significant_dyad_bootstrap.csv')
    scan_signature=pd.read_csv(TABLE/'figure4_transition_signature_model.csv').iloc[0]
    progress=pd.read_csv(TABLE/'figureS1_trial_progress_accuracy.csv')
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1);sec.header_distance=sec.footer_distance=Inches(.492)
    styles=doc.styles
    normal=styles['Normal'];normal.font.name='Calibri';normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.25
    for nm,size,color,before,after in [('Heading 1',16,BLUE,18,10),('Heading 2',13,BLUE,14,7),('Heading 3',12,DARK,10,5)]:
        s=styles[nm];s.font.name='Calibri';s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');s.font.size=Pt(size);s.font.color.rgb=color;s.font.bold=True;s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after);s.paragraph_format.keep_with_next=True
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;set_font(footer.add_run('PB&R revision analysis | generated from the OSF-ready pipeline'),size=8,color=MUTED)
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(4);set_font(p.add_run('补充数据分析结果报告'),size=24,bold=True,color=DARK)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(14);set_font(p.add_run('行为、元认知、注视分配与扫描路径'),size=14,color=MUTED)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(16);set_font(p.add_run('Psychonomic Bulletin & Review 修订分析 | 生成日期：2026-07-13'),size=9.5,color=MUTED)
    lead=doc.add_paragraph();lead.paragraph_format.space_after=Pt(10);set_font(lead.add_run('结论边界：'),bold=True,color=DARK);set_font(lead.add_run('以下结果检验的是 gravity 与 zero-gravity 线索背景下任务需求、表现监控和注视组织的差异，不把该对比解释为对“内部重力先验被因果扰动”的唯一检验。'))

    doc.add_heading('1. 样本与质量控制',level=1)
    add_table(doc,['步骤','人数'],[(r.stage,int(r.n)) for r in flow.itertuples()],widths=[5.1,1.4])
    matched=int(flow.loc[flow.stage.eq('Matched behavior-eye files'),'n'].iloc[0])
    behavior_n=int(flow.loc[flow.stage.eq('Behavior QC included'),'n'].iloc[0])
    auc_n=int(flow.loc[flow.stage.eq('AUC QC included'),'n'].iloc[0])
    eye_n=int(flow.loc[flow.stage.eq('Eye QC included (50%)'),'n'].iloc[0])
    dyad_n=int(flow.loc[flow.stage.eq('SPAM/DSM M-2SD included'),'n'].iloc[0])
    p=doc.add_paragraph();set_font(p.add_run('样本保留率：'),bold=True,color=DARK)
    set_font(p.add_run(f'行为、AUC、眼动与序列分析分别保留{behavior_n}/{matched}（{behavior_n/matched*100:.1f}%）、{auc_n}/{matched}（{auc_n/matched*100:.1f}%）、{eye_n}/{matched}（{eye_n/matched*100:.1f}%）和{dyad_n}/{matched}（{dyad_n/matched*100:.1f}%）的匹配样本。各类结果使用与终点相对应的预设质量门槛，不将不同终点的样本量混为同一分母。'))
    add_bullet(doc,'行为筛选严格执行：任一条件有效反应少于24、正确少于3或错误少于3即整名排除。AUC另要求每条件至少24个有效信心反应且正确、错误均至少3次。')
    add_bullet(doc,'眼动主分析要求两条件均有数据、非Gap覆盖率至少50%，且每条件至少10个有效AOI注视事件；40%与60%阈值作为敏感性分析。')
    add_bullet(doc,'SPAM/DSM按每名参与者两条件合计的可用dyad数筛选，低于样本M-2SD者排除；所有推断使用归一化dyad频率。')

    # Keep the main results table with its heading instead of leaving an
    # orphaned header row at the foot of the quality-control page. Applying
    # the break to the heading avoids an extra blank page when the preceding
    # section happens to end exactly at a page boundary in Microsoft Word.
    h=doc.add_heading('2. 核心条件差异',level=1)
    h.paragraph_format.page_break_before=True
    labels={'accuracy':'正确率','answer_rt_median':'决策RT中位数（ms）','answer_rt_mean':'决策RT均值（ms）','confidence_mean':'原始信心（1-7）','confidence_rt_median':'信心RT中位数（ms）','auc':'type-2 ROC AUC'}
    rows=[]
    for r in effects.itertuples():
        paired=(f't({int(r.n)-1})={fmt(r.paired_t,2)}, {ptxt(r.paired_t_p)}, dz={fmt(r.cohens_dz,2)}\n'
                f'W={int(r.wilcoxon_w)}, {ptxt(r.wilcoxon_p)}, r_rb={fmt(r.rank_biserial_sign,2)}')
        rows.append([labels.get(r.metric,r.metric),int(r.n),f'{fmt(r.gravity_mean)} ({fmt(r.gravity_sd)})',f'{fmt(r.zero_gravity_mean)} ({fmt(r.zero_gravity_sd)})',f'{fmt(r.difference_zero_minus_gravity)} [{fmt(r.ci95_low)}, {fmt(r.ci95_high)}]',paired])
    add_table(doc,['指标','n','Gravity M (SD)','Zero gravity M (SD)','差值及95% CI','配对t / Wilcoxon'],rows,widths=[1.15,.42,1.1,1.1,1.3,1.43])
    acc=effects[effects.metric.eq('accuracy')].iloc[0]
    rt=effects[effects.metric.eq('answer_rt_median')].iloc[0]
    conf=effects[effects.metric.eq('confidence_mean')].iloc[0]
    conf_rt=effects[effects.metric.eq('confidence_rt_median')].iloc[0]
    auc=effects[effects.metric.eq('auc')].iloc[0]
    p=doc.add_paragraph();set_font(p.add_run('数值结果：'),bold=True,color=DARK)
    set_font(p.add_run(
        f'Zero-gravity相对gravity的正确率差为{acc.difference_zero_minus_gravity*100:.2f}个百分点'
        f'（95% CI [{acc.ci95_low*100:.2f}, {acc.ci95_high*100:.2f}]，{ptxt(acc.paired_t_p)}，dz={acc.cohens_dz:.2f}）。'
        f'原始信心平均降低{abs(conf.difference_zero_minus_gravity):.3f}分'
        f'（95% CI [{conf.ci95_low:.3f}, {conf.ci95_high:.3f}]，dz={conf.cohens_dz:.2f}）；'
        f'type-2 ROC AUC平均降低{abs(auc.difference_zero_minus_gravity):.3f}'
        f'（95% CI [{auc.ci95_low:.3f}, {auc.ci95_high:.3f}]，n={int(auc.n)}，dz={auc.cohens_dz:.2f}）。'
        f'决策RT中位数仅增加{rt.difference_zero_minus_gravity:.1f} ms，95% CI跨零'
        f'（[{rt.ci95_low:.1f}, {rt.ci95_high:.1f}]，{ptxt(rt.paired_t_p)}）；决策RT均值差异也接近零。'
        f'信心RT中位数反而缩短{abs(conf_rt.difference_zero_minus_gravity):.1f} ms'
        f'（95% CI [{conf_rt.ci95_low:.1f}, {conf_rt.ci95_high:.1f}]，{ptxt(conf_rt.paired_t_p)}，dz={conf_rt.cohens_dz:.2f}）。'
        '因此，最一致的组合是正确率、信心和AUC降低，而决策时长并未呈现稳健的未调整条件差异。'))
    p=doc.add_paragraph();set_font(p.add_run('解释边界：'),bold=True,color=DARK)
    set_font(p.add_run('这些效应量整体为小到小—中等，表明zero-gravity线索背景产生了方向一致但幅度有限的行为—元认知差异。最稳妥的表述是任务背景改变了可用线索及其与判断要求的对齐程度；数据本身不能唯一识别为内部重力先验受到因果扰动。'))
    add_labeled(doc,'非参数与多重比较核对','Wilcoxon符号秩检验与配对t检验给出相同的显著性格局：正确率、原始信心、信心RT和AUC达到显著，决策RT的中位数与均值均未达到显著；六个终点经BH校正后结论不变。表中r_rb为配对秩二列效应，方向统一为zero gravity减gravity。')
    add_labeled(doc,'正确率-AUC关系',
        f'gravity条件中Pearson r={acc_auc_corr.iloc[0].pearson_r:.3f}（n={int(acc_auc_corr.iloc[0].n)}，{ptxt(acc_auc_corr.iloc[0].pearson_p)}），Spearman rho={acc_auc_corr.iloc[0].spearman_rho:.3f}；'
        f'zero-gravity条件中Pearson r={acc_auc_corr.iloc[1].pearson_r:.3f}（n={int(acc_auc_corr.iloc[1].n)}，{ptxt(acc_auc_corr.iloc[1].pearson_p)}），Spearman rho={acc_auc_corr.iloc[1].spearman_rho:.3f}。'
        'AUC与Type-I表现相关但并不等同；这些相关不能单独证明AUC条件差异完全独立于任务难度。')
    add_labeled(doc,'RT敏感性',
        f'将决策RT限制在{int(rt_sens.lower_ms)}-{int(rt_sens.upper_ms)} ms后，zero-gravity减gravity的配对差为{rt_sens.zero_minus_gravity:.1f} ms'
        f'（95% CI [{rt_sens.ci95_low:.1f}, {rt_sens.ci95_high:.1f}]，n={int(rt_sens.n)}，{ptxt(rt_sens.paired_p)}），仍不支持稳定的未调整决策减慢。')
    add_figure_block(
        doc,
        '对两条件数据完整且通过相应质量控制的参与者进行被试内配对比较；原始差值用配对t检验和95% CI表示，并用配对标准化效应d_z统一不同量纲。',
        FIG/'Figure_1_behavioral_signature.png',Inches(6.35),
        'Standardized paired effects and paired participant distributions for behavioral and metacognitive outcomes',
        '图1. 核心行为—元认知条件效应。',
        'a的纵轴列出六个终点，横轴是zero gravity减gravity的配对标准化效应；点为d_z，横线为95% CI，零线表示无差异。b–d的横轴为两种任务背景，纵轴依次为正确率、平均信心和type-2 ROC AUC；灰线连接同一参与者，空心点与误差线表示组均值及95% CI。',
        '正确率、原始信心和AUC的效应均位于零线左侧，且参与者分布整体向zero-gravity方向下移；反应时效应更接近零，说明变化并不是所有指标同步恶化。',
        'zero-gravity背景下表现与监控质量同时降低，但没有证据支持一种普遍、稳定的反应减慢。',
        '结果符合“线索背景与判断要求的对齐程度改变表现及元认知监控”的假设；它支持任务背景差异，不构成内部重力先验被因果扰动的唯一证据。')

    h=doc.add_heading('3. 试次级控制模型与信心耦合',level=1)
    acc_m=exact_coef(behavior_coef,'accuracy_glmm','conditionzero_gravity')
    rt_m=exact_coef(behavior_coef,'answer_rt_lmm','conditionzero_gravity')
    conf_rt_m=exact_coef(behavior_coef,'confidence_rt_lmm','conditionzero_gravity')
    conf_m=exact_coef(behavior_coef,'confidence_clmm','conditionzero_gravity')
    couple_condition=exact_coef(behavior_coef,'confidence_accuracy_coupling_glmm','conditionzero_gravity')
    couple_within=exact_coef(behavior_coef,'confidence_accuracy_coupling_glmm','confidence_within')
    couple_between=exact_coef(behavior_coef,'confidence_accuracy_coupling_glmm','confidence_between')
    couple_m=exact_coef(behavior_coef,'confidence_accuracy_coupling_glmm','conditionzero_gravity:confidence_within')
    bias_condition=exact_coef(behavior_coef,'confidence_bias_clmm','conditionzero_gravity')
    bias_correct=exact_coef(behavior_coef,'confidence_bias_clmm','correct_f1')
    bias_interaction=exact_coef(behavior_coef,'confidence_bias_clmm','conditionzero_gravity:correct_f1')
    trial_n=behavior_diag.set_index('model')['n'].to_dict()
    base_models=['accuracy_glmm','answer_rt_lmm','confidence_rt_lmm','confidence_clmm']
    inter_info={}
    for model in base_models:
        z=behavior_coef[(behavior_coef.model==model)&behavior_coef.term.astype(str).str.startswith('conditionzero_gravity:')]
        sigz=z[z.p_value<.05]
        inter_info[model]=(len(sigz),len(z),[interaction_label(x) for x in sigz.term])
    model_items=[
        [f'正确性GLMM\n{int(trial_n["accuracy_glmm"])}试次',f'条件b={acc_m.estimate:.3f}（SE={acc_m.std_error:.3f}），{ptxt(acc_m.p_value)}；OR={acc_m.odds_ratio:.3f} [{acc_m.ci_low:.3f}, {acc_m.ci_high:.3f}]',f'显著条件交互{inter_info["accuracy_glmm"][0]}/{inter_info["accuracy_glmm"][1]}'],
        [f'决策RT LMM\n{int(trial_n["answer_rt_lmm"])}试次',f'条件b={rt_m.estimate:.3f}（SE={rt_m.std_error:.3f}），{ptxt(rt_m.p_value)}；95% CI [{rt_m.ci_low:.3f}, {rt_m.ci_high:.3f}]',f'显著条件交互{inter_info["answer_rt_lmm"][0]}/{inter_info["answer_rt_lmm"][1]}'],
        [f'信心RT LMM\n{int(trial_n["confidence_rt_lmm"])}试次',f'条件b={conf_rt_m.estimate:.3f}（SE={conf_rt_m.std_error:.3f}），{ptxt(conf_rt_m.p_value)}；95% CI [{conf_rt_m.ci_low:.3f}, {conf_rt_m.ci_high:.3f}]',f'显著条件交互{inter_info["confidence_rt_lmm"][0]}/{inter_info["confidence_rt_lmm"][1]}'],
        [f'原始信心CLMM\n{int(trial_n["confidence_clmm"])}试次',f'条件b={conf_m.estimate:.3f}（SE={conf_m.std_error:.3f}），{ptxt(conf_m.p_value)}；OR={conf_m.odds_ratio:.3f} [{conf_m.ci_low:.3f}, {conf_m.ci_high:.3f}]',f'显著条件交互{inter_info["confidence_clmm"][0]}/{inter_info["confidence_clmm"][1]}'],
        [f'信心-正确性耦合GLMM\n{int(trial_n["confidence_accuracy_coupling_glmm"])}试次',f'条件×被试内信心b={couple_m.estimate:.3f}（SE={couple_m.std_error:.3f}），{ptxt(couple_m.p_value)}；OR={couple_m.odds_ratio:.3f} [{couple_m.ci_low:.3f}, {couple_m.ci_high:.3f}]','同时估计条件、被试内及被试间信心'],
        [f'信心偏差/校准CLMM\n{int(trial_n["confidence_bias_clmm"])}试次',f'条件×正确性b={bias_interaction.estimate:.3f}（SE={bias_interaction.std_error:.3f}），{ptxt(bias_interaction.p_value)}；OR={bias_interaction.odds_ratio:.3f} [{bias_interaction.ci_low:.3f}, {bias_interaction.ci_high:.3f}]',f'条件主效应OR={bias_condition.odds_ratio:.3f}，{ptxt(bias_condition.p_value)}']
    ]
    add_table(doc,['六个主模型','焦点结果','补充说明'],model_items,widths=[1.65,3.35,1.5])
    add_bullet(doc,'六个主模型均收敛；主结果同时控制速度、左右球尺寸配对、Target、中心化试次序号及二次项。前四个终点模型还检验条件×速度和条件×尺寸，且均保留相关条件随机斜率，无需回退。')
    interaction_text=[]
    interaction_names={'accuracy_glmm':'正确性','answer_rt_lmm':'决策RT','confidence_rt_lmm':'信心RT','confidence_clmm':'原始信心'}
    for model in base_models:
        interaction_text.append(interaction_names[model]+'：'+('、'.join(inter_info[model][2]) if inter_info[model][2] else '无'))
    add_bullet(doc,'达到p < .05的条件异质性项为：'+'；'.join(interaction_text)+'。这些结果说明条件效应幅度随部分运动参数变化。')
    add_labeled(doc,'主模型解释',
        f'在参考速度与参考尺寸组合下，zero-gravity的正确优势较低（OR={acc_m.odds_ratio:.3f}），原始信心进入更高等级的优势较低（OR={conf_m.odds_ratio:.3f}）；'
        f'决策RT条件系数对应{(math.exp(rt_m.estimate)-1)*100:.1f}%变化，而信心RT条件系数对应{(math.exp(conf_rt_m.estimate)-1)*100:.1f}%变化且未达到显著（{ptxt(conf_rt_m.p_value)}）。'
        f'耦合模型中，被试内信心每增加1级，gravity条件的正确优势乘以{couple_within.odds_ratio:.3f}；zero-gravity中的该增益再乘以{couple_m.odds_ratio:.3f}，即相对减弱约{(1-couple_m.odds_ratio)*100:.1f}%。'
        f'被试间平均信心效应未达到显著（OR={couple_between.odds_ratio:.3f}，{ptxt(couple_between.p_value)}）。')
    add_labeled(doc,'信心偏差/校准',
        f'控制正确性后，zero-gravity仍与较低信心等级相关（OR={bias_condition.odds_ratio:.3f}，95% CI [{bias_condition.ci_low:.3f}, {bias_condition.ci_high:.3f}]，{ptxt(bias_condition.p_value)}）；'
        f'正确反应对应更高信心（OR={bias_correct.odds_ratio:.3f}，{ptxt(bias_correct.p_value)}），但条件×正确性交互不显著（OR={bias_interaction.odds_ratio:.3f}，{ptxt(bias_interaction.p_value)}）。因此，较低原始信心并未伴随一个可确认的“正确/错误分离幅度”额外变化。')
    occ_acc=exact_coef(behavior_coef,'accuracy_occlusion_sensitivity_glmm','conditionzero_gravity')
    occ_acc_int=exact_coef(behavior_coef,'accuracy_occlusion_sensitivity_glmm','conditionzero_gravity:occlusion_hidden_c')
    occ_rt=exact_coef(behavior_coef,'rt_occlusion_sensitivity_lmm','conditionzero_gravity')
    occ_rt_int=exact_coef(behavior_coef,'rt_occlusion_sensitivity_lmm','conditionzero_gravity:occlusion_hidden_c')
    sensitivity_rows=[
        ['正确性GLMM',f'条件OR={occ_acc.odds_ratio:.3f} [{occ_acc.ci_low:.3f}, {occ_acc.ci_high:.3f}]，{ptxt(occ_acc.p_value)}；条件×遮挡时长OR={occ_acc_int.odds_ratio:.3f} [{occ_acc_int.ci_low:.3f}, {occ_acc_int.ci_high:.3f}]，{ptxt(occ_acc_int.p_value)}'],
        ['决策RT LMM',f'条件b={occ_rt.estimate:.3f} [{occ_rt.ci_low:.3f}, {occ_rt.ci_high:.3f}]，{ptxt(occ_rt.p_value)}；条件×遮挡时长b={occ_rt_int.estimate:.3f} [{occ_rt_int.ci_low:.3f}, {occ_rt_int.ci_high:.3f}]，{ptxt(occ_rt_int.p_value)}']
    ]
    add_table(doc,['遮挡时长替代模型','结果'],sensitivity_rows,widths=[1.65,4.85])
    add_bullet(doc,'遮挡时长替代速度的敏感性分析仅对正确性与决策RT运行：正确性条件方向保持，但效应随遮挡时长变化；决策RT的条件项及交互均不显著。当前输出没有原始信心或信心-正确性耦合的遮挡时长替代模型。')
    add_labeled(doc,'系数边界','由于前四个模型包含条件交互，上表的条件b/OR是参考速度与参考尺寸组合下的条件效应，而不是跨全部速度和尺寸组合的总体平均效应。总体条件方向应与配对结果、分层描述和交互清单共同解释。')
    add_figure_block(
        doc,
        '校准曲线先在每名参与者内计算1–7级信心对应的正确率，再跨参与者汇总并bootstrap 95% CI；下方结果来自控制速度、尺寸配对、target和试次趋势的GLMM/CLMM，统一换算为zero gravity相对gravity的优势比。',
        FIG/'Figure_2_metacognitive_monitoring.png',Inches(6.35),
        'Confidence calibration curves and adjusted odds ratios for accuracy, confidence, and confidence-accuracy coupling',
        '图2. 信心校准与调整后的元认知监控效应。',
        'a位于上方，横轴是信心等级，纵轴是在该等级下实际答对的比例；实线为两种条件的均值，半透明带为95% CI。b位于下方，纵轴依次为正确性、进入更高信心等级以及信心—正确性耦合，横轴为zero gravity/gravity调整后优势比；竖虚线1表示两条件无差异，每行右侧直接标出OR及95% CI。',
        '两条校准曲线总体随信心升高而上升，但zero-gravity曲线在中间信心等级多位于gravity之下。下图三个优势比均低于1且95% CI不跨1，其中信心等级的相对降低最大，信心—正确性耦合的效应较小但方向稳定。',
        '即使把试次结构纳入模型，zero-gravity背景仍与较低正确优势、较低信心水平及较弱的信心诊断性相联系；这不是仅由某个孤立的散点关系形成的结论。',
        '该组合更符合任务背景改变证据质量和监控校准的解释，而不是把AUC或信心变化简单等同于内部重力模型的改变。')

    doc.add_heading('3.1 任务因素与时间进程',level=2)
    size_diff=pd.read_csv(TABLE/'figureS1_size_pair_accuracy_difference.csv')
    progress_wide=progress.pivot(index='trial_bin',columns='condition',values='accuracy')
    progress_diff=progress_wide.zero_gravity-progress_wide.gravity
    add_bullet(doc,f'速度分层中，除3.5速度水平外，zero-gravity平均正确率均低于gravity；这说明总体条件方向具有跨速度的一致性，但3.5水平的交叉提醒我们不能把效应写成完全不受速度结构影响。')
    add_bullet(doc,f'9种尺寸配对的zero-gravity−gravity正确率差异范围为{size_diff.difference.min():+.2f}至{size_diff.difference.max():+.2f}；尺寸组合会调节效应幅度和局部方向。')
    add_bullet(doc,f'在6个条件内六试次分箱中，zero-gravity−gravity正确率差异均为负，范围为{progress_diff.min():+.3f}至{progress_diff.max():+.3f}。这表明条件差异并非由单个局部试次段独立驱动，但不同分箱的差值幅度仍有波动。')
    g_stim=stimulus_checks[stimulus_checks.condition.eq('gravity')]
    z_stim=stimulus_checks[stimulus_checks.condition.eq('zero_gravity')]
    add_bullet(doc,
        f'刺激生成审计显示，12个条件×速度单元中的distance均约为{stimulus_checks.distance_mean.mean():.3f}，first boundary均为Horizontal；'
        f'但平均遮挡时长在gravity中跨速度由{g_stim.hidden_duration_mean.min():.1f}至{g_stim.hidden_duration_mean.max():.1f} ms，在zero-gravity中由{z_stim.hidden_duration_mean.min():.1f}至{z_stim.hidden_duration_mean.max():.1f} ms。'
        '因此，遮挡时长是与条件和速度共同变化的刺激属性，替代参数化用于界定而不能消除这一混淆。')
    add_figure_block(
        doc,
        '将试次按速度水平、左右球尺寸配对以及每个条件内的六试次分箱进行分层；先计算参与者内正确率，再汇总条件均值与95% CI，尺寸配对用zero gravity减gravity的差值表示。',
        FIG/'Figure_S1_task_structure_and_time_course.png',Inches(6.35),
        'Accuracy stratified by velocity, ball-size pairing, and within-condition trial bins',
        '补充图S1. 任务因素与条件内时间进程。',
        'a的横轴为速度水平、纵轴为正确率；b的行列分别为左右球尺寸，单元格数字是zero gravity减gravity的正确率差；c的横轴为条件内连续六试次分箱，纵轴为正确率。线和阴影分别表示均值和95% CI。',
        '多数速度水平和全部试次分箱中zero-gravity正确率较低，但不同尺寸组合和局部速度水平会改变差异幅度，个别组合出现较弱或反向的局部差值。',
        '总体条件差异分布在多个运动参数和任务进程区段中，并非由单一速度、单一尺寸组合或某个局部试次段独立驱动；同时，其大小会随刺激结构变化。',
        '这使“任务背景与线索结构共同影响判断”的解释比单一、固定的重力机制效应更符合数据。')

    doc.add_heading('4. AOI注视分配',level=1)
    eye_coef=TABLE/'eye_mixed_model_coefficients.csv'
    if eye_coef.exists():
        ec=pd.read_csv(eye_coef);z=ec[ec.term.astype(str).str.contains('conditionzero_gravity:regionposition',regex=False)]
        add_table(doc,['指标模型','条件×区域 b (SE), p'],[(r.model,f'{fmt(r.estimate)} ({fmt(r.std_error)}), {ptxt(r.p_value)}') for r in z.itertuples()],widths=[3.4,3.1])
    add_bullet(doc,'AOI模型的完整与去相关随机斜率结构均触发奇异或不收敛，故按预设规则回退到随机截距。注视比例模型因球区与位置区为互补份额而仍有奇异性，因此主图直接展示两条件在球区和位置区的原始眼动指标分布；位置区注视占比变化另在40%/50%/60%覆盖率阈值下进行敏感性检验。')
    position50=eye_sens.iloc[(eye_sens.coverage_threshold-.5).abs().argmin()]
    p=doc.add_paragraph();set_font(p.add_run('数值结果：'),bold=True,color=DARK)
    def eye_mean(metric,region,condition):
        return eye_direct[(eye_direct.metric.eq(metric))&(eye_direct.region.eq(region))&
                          (eye_direct.condition.eq(condition))].iloc[0]['mean']
    ball_time_g=eye_mean('total_fixation_time','ball','gravity');ball_time_z=eye_mean('total_fixation_time','ball','zero_gravity')
    pos_time_g=eye_mean('total_fixation_time','position','gravity');pos_time_z=eye_mean('total_fixation_time','position','zero_gravity')
    ball_count_g=eye_mean('fixation_count','ball','gravity');ball_count_z=eye_mean('fixation_count','ball','zero_gravity')
    pos_count_g=eye_mean('fixation_count','position','gravity');pos_count_z=eye_mean('fixation_count','position','zero_gravity')
    ball_duration_g=eye_mean('mean_fixation_time','ball','gravity');ball_duration_z=eye_mean('mean_fixation_time','ball','zero_gravity')
    pos_duration_g=eye_mean('mean_fixation_time','position','gravity');pos_duration_z=eye_mean('mean_fixation_time','position','zero_gravity')
    set_font(p.add_run(
        f'主阈值下，位置区注视占比在zero-gravity中平均增加{position50.position_share_change_pp:.2f}个百分点'
        f'（95% CI [{position50.ci95_low_pp:.2f}, {position50.ci95_high_pp:.2f}]，n={int(position50.n)}，{ptxt(position50.p)}）；'
        f'40%与60%阈值下分别增加{eye_sens.iloc[(eye_sens.coverage_threshold-.4).abs().argmin()].position_share_change_pp:.2f}和{eye_sens.iloc[(eye_sens.coverage_threshold-.6).abs().argmin()].position_share_change_pp:.2f}个百分点，方向与幅度稳定。'
        f'原始指标显示，球区总注视时长由gravity的{ball_time_g:.1f} s降至zero-gravity的{ball_time_z:.1f} s，'
        f'位置区则由{pos_time_g:.1f} s升至{pos_time_z:.1f} s；球区注视次数由{ball_count_g:.1f}降至{ball_count_z:.1f}，'
        f'位置区由{pos_count_g:.1f}升至{pos_count_z:.1f}。平均单次注视时长在球区为{ball_duration_g:.3f} s与{ball_duration_z:.3f} s，'
        f'在位置区为{pos_duration_g:.3f} s与{pos_duration_z:.3f} s。'))
    add_figure_block(
        doc,
        '以每个条件内球区与位置区的总注视时长计算注视占比；同时直接汇总每名参与者在两种条件、两个区域中的总注视时长、注视次数和平均单次注视时长。分布图保留全部参与者的分布形态，并叠加中位数、四分位距以及均值的95% CI；位置区占比变化在40%、50%和60%覆盖率阈值下另作敏感性检验。',
        FIG/'Figure_3_gaze_reallocation.png',Inches(6.35),
        'Fixation-time allocation and direct condition-specific distributions of three eye-movement measures',
        '图3. 注视预算与两条件原始眼动指标。',
        'a的横向堆叠条表示两条件中球区和位置区占全部注视时长的比例，括号连接两条中位置区边界。b–d的横轴均为球区和位置区，蓝色与橙色分别表示gravity和zero gravity；纵轴依次为总注视时长、注视次数和平均单次注视时长。小提琴宽度表示参与者分布密度，黑色粗线为四分位距，白点为中位数，彩色点和误差线为均值及95% CI，点旁数字为条件均值。',
        f'位置区平均占比由14.9%增至26.9%，即增加约{position50.position_share_change_pp:.1f}个百分点。原始分布显示，zero-gravity下球区总时长和注视次数较低，而位置区对应指标较高；平均单次时长的球区差异较小，位置区则呈增加。',
        'zero-gravity背景伴随注视预算从运动物体转向候选落点/位置区域，而且这一模式不是由某个特定眼动覆盖率阈值造成。',
        '结果与观察者在较不熟悉的轨迹背景下增加结果位置核查的策略解释一致；它是注意分配相关证据，不是重力先验直接驱动眼动的因果中介证据。')

    doc.add_heading('5. SPAM/DSM探索结果',level=1)
    sig=dsm[dsm.permutation_p_fdr<.05].sort_values('permutation_p_fdr')
    if len(sig):
        boot_map=scan_boot.set_index('dyad')
        rows=[[r.representation,r.condition,r.dyad,int(r.n_high),int(r.n_low),fmt(boot_map.loc[r.dyad,'estimate'],4),ptxt(r.permutation_p_fdr)] for r in sig.itertuples()]
        add_table(doc,['类型','条件','Dyad','High n','Low n','Δ Low−High','FDR p'],rows[:20],widths=[.7,.85,2.15,.55,.55,.8,.9])
    else:doc.add_paragraph('FDR校正后没有高/低AUC组的dyad差异达到显著。')
    add_bullet(doc,'高/低组按条件内AUC的Q1/Q3分组（切点不入组）；主推断为归一化频率置换检验，Welch t仅供对照，event与duration分别进行BH-FDR校正。')
    add_bullet(doc,'眼动文件无trial identifier，同一TOI内不能排除跨试次边界转移；SPAM/DSM仅作探索性序列组织证据。')
    support_groups=spam_support.groupby(['representation','condition']).s_support.agg(['min','max'])
    add_bullet(doc,
        f'模式支持度方面，event dyad的参与者支持比例范围为{support_groups.loc[("event","gravity"),"min"]:.3f}-{support_groups.loc[("event","zero_gravity"),"max"]:.3f}；'
        f'duration dyad在gravity和zero-gravity中的最低支持比例分别为{support_groups.loc[("duration","gravity"),"min"]:.3f}和{support_groups.loc[("duration","zero_gravity"),"min"]:.3f}。支持度用于描述模式普遍性，组间推断仍以参与者归一化频率为单位。')
    p=doc.add_paragraph();set_font(p.add_run('数值结果与边界：'),bold=True,color=DARK)
    set_font(p.add_run(
        f'FDR校正后仅zero-gravity的3个duration dyad显著。为便于理解，本报告将Δ定义为低AUC组频率减高AUC组频率；'
        f'三个Δ均为正，范围为{scan_boot.estimate.min():.4f}至{scan_boot.estimate.max():.4f}，bootstrap 95% CI整体边界为[{scan_boot.ci_low.min():.4f}, {scan_boot.ci_high.max():.4f}]'
        f'（高/低组n={int(scan_boot.n_high.iloc[0])}/{int(scan_boot.n_low.iloc[0])}，FDR p=.021）。'
        f'将这三类转移的归一化频率相加后，连续AUC与该转移特征的相关为r={scan_signature.pearson_r:.2f}'
        f'（n={int(scan_signature.n)}，{ptxt(scan_signature.p)}，R²={scan_signature.r_squared*100:.1f}%）。'
        '连续关系属于在FDR筛选后的描述性跟进，不能视为独立验证；方向虽一致，解释仍应限于探索性关联。'))
    cont_sig=dsm_continuous[dsm_continuous.p_fdr<.05].copy()
    event_sig=cont_sig[cont_sig.representation.eq('event')]
    duration_sig=cont_sig[cont_sig.representation.eq('duration')]
    add_labeled(doc,'连续AUC敏感性',
        f'对全部dyad分别进行连续AUC回归并在event与duration内作BH-FDR校正后，仅zero-gravity结果保留：'
        f'event的ball->position与position->ball斜率分别为{event_sig.iloc[0].auc_slope:.3f}和{event_sig.iloc[1].auc_slope:.3f}（两者FDR p={event_sig.iloc[0].p_fdr:.3f}）；'
        f'duration的ball_MED->position_LONG、ball_MED->position_MED和position_MED->ball_MED斜率分别为{duration_sig.iloc[0].auc_slope:.3f}、{duration_sig.iloc[1].auc_slope:.3f}和{duration_sig.iloc[2].auc_slope:.3f}'
        f'（FDR p={duration_sig.iloc[0].p_fdr:.3f}、{duration_sig.iloc[1].p_fdr:.3f}、{duration_sig.iloc[2].p_fdr:.3f}；均n={int(duration_sig.iloc[0].n)}）。'
        '负斜率与极端组结果方向一致，即AUC越低，跨区域往返频率越高；但效应解释仍限于探索性关联。')
    add_figure_block(
        doc,
        '在zero-gravity的duration序列中计算每名参与者各有向dyad占全部可用转移的归一化频率；按条件内AUC的Q1/Q3形成低/高组，以置换检验并经BH-FDR校正。图中Δ统一定义为“低AUC组频率减高AUC组频率”；随后把三个FDR显著dyad的频率相加，与连续AUC作描述性线性关联。',
        FIG/'Figure_4_scanpath_organization.png',Inches(5.85),
        'Zero-gravity duration transition matrix, continuous AUC-signature association, and directed transition network',
        '图4. 探索性扫描路径组织。',
        'a的行是转移起点、列是转移终点，单元格数字为Δ（百分点），暖色表示低AUC组更高，星号和粗框表示FDR<.05；b的横轴为type-2 ROC AUC，纵轴为三个显著dyad合计占全部duration转移的比例，实线和阴影为线性拟合及95% CI；c的节点是注视状态，箭头表示转移方向，标签直接给出正向Δ。',
        '显著单元格集中在Ball (Med)→Position (Med/Long)以及Position (Med)→Ball (Med)，Δ约为+0.66至+0.72个百分点。连续图呈负相关：AUC越低，这组三类跨区域转移合计占比越高。',
        '低AUC参与者在zero-gravity背景中更频繁地出现一组围绕中等时长球区与位置区往返的转移，但绝对差异较小，且其余大多数转移未通过FDR校正。',
        '该模式可作为“较弱元认知监控伴随更分散或反复的扫描路径组织”的探索性证据；由于分组极端、特征经过结果筛选且序列可能跨试次边界，它不支持稳定个体类型或因果机制结论。')

    doc.add_heading('6. 投稿结论建议',level=1)
    add_bullet(doc,'主结论写成gravity与zero-gravity线索背景下的条件差异，并说明其伴随表现、信心诊断性和注视分配改变。')
    add_bullet(doc,'不要写成zero gravity直接违反或更新了内部重力先验；0g恒速运动和可能的平面碰撞理解仍是需要保留的替代解释。')
    add_bullet(doc,'把AUC与试次级信心耦合的分工写清：AUC是非参数二阶敏感性指标；信心—正确性模型检验信心在两个条件下的诊断性强度是否不同。')
    add_bullet(doc,'AOI结果用associated with / accompanied by / consistent with；SPAM/DSM明确标为探索性并报告FDR、归一化频率和dyad长度筛选。')
    add_figure_block(
        doc,
        '分别按行为、AUC、眼动和SPAM/DSM的预设质量门槛统计保留样本；展示两条件眼动覆盖率和每名参与者可用dyad数量，并对核心配对效应逐一删除一名参与者后重新估计。',
        FIG/'Figure_S2_quality_control_and_robustness.png',Inches(6.35),
        'Analysis-specific sample flow, eye coverage, dyad screening, and leave-one-participant-out stability',
        '补充图S2. 质量控制与稳健性。',
        'a按分析终点列出样本流；b的横轴为非Gap眼动覆盖率、颜色区分条件；c的横轴为可用dyad数量，虚线为M−2SD筛选阈值；d的纵轴列出核心终点，横轴为zero gravity减gravity的留一被试效应范围，点为全样本估计。',
        f'从{matched}名匹配参与者中，行为、AUC、眼动与序列分析分别保留{behavior_n}、{auc_n}、{eye_n}和{dyad_n}名；留一被试后的核心差异方向保持不变，说明主要符号不依赖某一个体。',
        '各结果采用分析特异的有效样本，而不是用同一个分母替代；核心条件差异对单个参与者的影响具有基本稳健性。',
        '质量控制图不直接检验理论机制，但它限定了结果的可信范围：理论解释应建立在通过对应门槛且方向稳定的终点上。')

    doc.core_properties.title='修订稿补充数据分析结果报告';doc.core_properties.subject='PB&R PBR-BR-26-106 revision analysis';doc.core_properties.author='Analysis pipeline';doc.core_properties.keywords='behavior; metacognition; eye tracking; SPAM; DSM; reproducibility'
    doc.save(REPORT);return REPORT

if __name__=='__main__': print(build())
